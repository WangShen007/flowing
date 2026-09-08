import asyncio
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pptx import Presentation

from app.api.routes import ai_ppt
from app.api.routes.auth import AccountInfo
from app.core.storage import SQLiteStore
from app.skills.ai_ppt import skill as ai_ppt_skill
from app.skills.ai_ppt.skill import AIPPTSkill
from app.skills.base import SkillContext


@pytest.fixture
def isolated_ai_ppt(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    database = SQLiteStore(tmp_path / "ai-ppt.sqlite3")
    base_dir = tmp_path / "ai-ppt"
    source_dir = base_dir / "sources"
    template_dir = base_dir / "templates"
    preview_dir = base_dir / "previews"
    builtin_dir = tmp_path / "builtin"
    for path in (source_dir, template_dir, preview_dir, builtin_dir):
        path.mkdir(parents=True)

    monkeypatch.setattr(ai_ppt, "store", database)
    monkeypatch.setattr(ai_ppt, "ai_ppt_base_dir", lambda: base_dir)
    monkeypatch.setattr(ai_ppt, "_source_dir", lambda: source_dir)
    monkeypatch.setattr(ai_ppt, "_template_dir", lambda: template_dir)
    monkeypatch.setattr(ai_ppt, "_preview_dir", lambda: preview_dir)
    monkeypatch.setattr(ai_ppt, "_builtin_template_dir", lambda: builtin_dir)
    monkeypatch.setattr(ai_ppt_skill, "ai_ppt_base_dir", lambda: base_dir)

    owned = source_dir / "owned.pptx"
    _write_valid_ppt(owned)
    legacy = source_dir / "legacy.pptx"
    _write_valid_ppt(legacy)
    for item in ai_ppt._default_templates():
        _write_valid_ppt(builtin_dir / item["filename"])
    (builtin_dir / "not-public.pptx").write_bytes(b"private")
    ai_ppt.register_ai_ppt_file("owned.pptx", "alice", "source", file_id="owned")

    account = {"value": AccountInfo(account="alice", name="Alice")}
    app = FastAPI()
    app.include_router(ai_ppt.router, prefix="/api/v1")
    app.dependency_overrides[ai_ppt.get_current_account] = lambda: account["value"]
    app.dependency_overrides[ai_ppt.get_current_account_optional] = lambda: account["value"]
    with TestClient(app) as client:
        yield client, account, database, base_dir, source_dir, template_dir, preview_dir, builtin_dir


def test_file_download_preview_and_action_are_owner_bound(isolated_ai_ppt, monkeypatch):
    client, account, _database, _base, _source, _templates, _previews, _builtin = isolated_ai_ppt

    assert client.get("/api/v1/ai-ppt/files/owned.pptx").status_code == 200
    account["value"] = AccountInfo(account="bob", name="Bob")
    assert client.get("/api/v1/ai-ppt/files/owned.pptx").status_code == 404
    account["value"] = None
    assert client.get("/api/v1/ai-ppt/files/owned.pptx").status_code == 404

    # Builtin templates are public only through the explicit default-template allowlist.
    assert client.get("/api/v1/ai-ppt/files/ppt_master_dark_tech.pptx").status_code == 200
    assert client.get("/api/v1/ai-ppt/files/not-public.pptx").status_code == 404

    account["value"] = AccountInfo(account="alice", name="Alice")
    preview = client.get("/api/v1/ai-ppt/files/owned.pptx/preview")
    assert preview.status_code == 200
    preview_id = preview.json()["data"]["slides"][0]["image_url"].split("/")[-2]
    assert client.get(preview.json()["data"]["slides"][0]["image_url"]).status_code == 200
    account["value"] = AccountInfo(account="bob", name="Bob")
    assert client.get(f"/api/v1/ai-ppt/previews/{preview_id}/slide_001.png").status_code == 404
    account["value"] = None
    assert client.get(f"/api/v1/ai-ppt/previews/{preview_id}/slide_001.png").status_code == 404

    builtin_preview = client.get("/api/v1/ai-ppt/files/ppt_master_dark_tech.pptx/preview")
    assert builtin_preview.status_code == 200
    builtin_image = builtin_preview.json()["data"]["slides"][0]["image_url"]
    assert client.get(builtin_image).status_code == 200

    async def fake_action(*_args, **_kwargs):
        return {"success": True, "message": "ok", "metadata": {}}

    monkeypatch.setattr(ai_ppt, "_run_feishu_action", fake_action)
    account["value"] = AccountInfo(account="alice", name="Alice")
    assert client.post("/api/v1/ai-ppt/actions", json={"filename": "owned.pptx", "action": "upload"}).status_code == 200
    account["value"] = AccountInfo(account="bob", name="Bob")
    assert client.post("/api/v1/ai-ppt/actions", json={"filename": "owned.pptx", "action": "upload"}).status_code == 404


def test_legacy_files_and_model_source_are_denied_but_builtin_is_allowed(isolated_ai_ppt):
    _client, _account, _database, _base, _source, _templates, _previews, _builtin = isolated_ai_ppt
    assert ai_ppt.resolve_ai_ppt_file("legacy.pptx", "alice") is None
    assert ai_ppt.resolve_ai_ppt_file("owned", "bob") is None
    assert ai_ppt.resolve_ai_ppt_file("ppt_master_dark_tech", "bob") is not None
    assert ai_ppt.resolve_ai_ppt_file("not-public.pptx", "alice") is None

    skill = AIPPTSkill()
    context = SkillContext(session_id="session", user_id="bob", message="修改 PPT")
    events = asyncio.run(_collect(skill.execute_stream(context, source_ppt_id="owned")))
    assert any(event.get("type") == "metadata" and event["data"].get("ai_ppt_forbidden") for event in events)

    generated = skill._render_deck(
        outline={
            "title": "归属测试",
            "subtitle": "",
            "slides": [{"title": "封面", "bullets": [], "speaker_notes": "", "layout": "cover"}],
        },
        query="生成测试 PPT",
        style="商务",
        owner_account="alice",
    )
    assert ai_ppt.resolve_ai_ppt_file(generated["filename"], "alice") is not None
    assert ai_ppt.resolve_ai_ppt_file(generated["filename"], "bob") is None


async def _collect(iterator):
    return [event async for event in iterator]


def _write_valid_ppt(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 1000, 1000).text = "isolation"
    presentation.save(path)
